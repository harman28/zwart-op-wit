import sys
import openpyxl

path = sys.argv[1]
sheet = sys.argv[2]
max_rows = int(sys.argv[3]) if len(sys.argv) > 3 else 15
max_cols = int(sys.argv[4]) if len(sys.argv) > 4 else 20

wb = openpyxl.load_workbook(path, data_only=True)
ws = wb[sheet]
print(f"sheet={sheet} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
for r in range(1, min(max_rows, ws.max_row) + 1):
    row = []
    for c in range(1, min(max_cols, ws.max_column) + 1):
        v = ws.cell(row=r, column=c).value
        row.append("" if v is None else str(v))
    print(f"R{r}: " + " | ".join(row))
