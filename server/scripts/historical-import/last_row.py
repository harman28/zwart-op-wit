import sys
import openpyxl

path = sys.argv[1]
sheet = sys.argv[2]

wb = openpyxl.load_workbook(path, data_only=True)
ws = wb[sheet]
last = 0
for r in range(1, ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        if ws.cell(row=r, column=c).value not in (None, ""):
            last = r
print(f"sheet={sheet} last_nonempty_row={last} of max_row={ws.max_row}, max_col={ws.max_column}")
