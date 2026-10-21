import sys
import openpyxl

path = sys.argv[1]
wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
print("Sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n=== {name} === max_row={ws.max_row} max_col={ws.max_column}")
