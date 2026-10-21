import sys
import re
import openpyxl

path = sys.argv[1]
sheets = ["S4 Matchups", "S5 Matchups", "S6 Matchups", "S7 Matchups", "S8 Matchups"]

wb = openpyxl.load_workbook(path, data_only=True)

odd_re = re.compile(r"odd", re.I)
rep_re = re.compile(r"rep", re.I)
r_re = re.compile(r"^\d+R$", re.I)

for sheet in sheets:
    ws = wb[sheet]
    print(f"\n--- {sheet} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    odd_examples = []
    rep_examples = []
    r_examples = []
    x_count = 0
    non_x_non_blank = set()
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            s = str(v)
            if odd_re.search(s) and len(odd_examples) < 6:
                odd_examples.append((r, c, s))
            if rep_re.search(s) and len(rep_examples) < 6:
                rep_examples.append((r, c, s))
            if r_re.match(s.strip()) and len(r_examples) < 6:
                r_examples.append((r, c, s))
            if s.strip() == "X":
                x_count += 1
            elif s.strip() not in ("",) and not re.match(r"^-?\d+(\.\d+)?$", s.strip()):
                non_x_non_blank.add(s.strip())
    print("odd examples:", odd_examples)
    print("rep examples:", rep_examples)
    print("R-suffix score examples:", r_examples)
    print("X count:", x_count)
